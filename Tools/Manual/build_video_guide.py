"""Builds the FeelKit Video Guide as a Word document, and as a PDF with --pdf.

The source is the part "## Video production guide" at the end of Docs/FeelKit_3_Manual_Plan.md. Its headings start one
level down (### is a chapter, #### a section, ##### a subsection); everything else uses the manual's Markdown subset,
parsed and styled by build_manual.py, which this script imports and does not change. Tables whose first column is
"Shot" are shot-by-shot scripts and go on landscape pages.

The cover is drawn into Docs/Manual/design/video_cover.png from the recipe editor picture that the manual's cover uses.
The build prints, for each script, the narration's word count against the time the shots allow, and checks the text
against the manual's writing rules (lint_manual.py).

Run:

    python build_video_guide.py                writes Docs/Manual/out/FeelKit_Video_Guide.docx
    python build_video_guide.py --pdf          also updates the fields in Word and exports the PDF
    python build_video_guide.py --release      leaves out the draft line on the title page
    python build_video_guide.py --source F     reads the guide from file F instead of the Manual Plan
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

import build_manual as bm
from manual_tables import MANUAL, PLAN

GUIDE_HEADING = "## Video production guide"
TITLE = "FeelKit Video Guide"
DOCX = MANUAL / "out" / "FeelKit_Video_Guide.docx"
COVER = MANUAL / "design" / "video_cover.png"
HERO = Path("B:/NewUE5Project/GameFeelDev/Saved/FeelKit/Manual/Editor_HeavyHit.png")

PORTRAIT = dict(width=21.0, height=29.7, left=2.7, right=2.7, top=2.6, bottom=2.4, header=1.2, footer=1.1)
LANDSCAPE = dict(width=29.7, height=21.0, left=2.2, right=2.2, top=1.7, bottom=1.5, header=0.8, footer=0.7)
PORTRAIT_TEXT = PORTRAIT["width"] - PORTRAIT["left"] - PORTRAIT["right"]
LANDSCAPE_TEXT = LANDSCAPE["width"] - LANDSCAPE["left"] - LANDSCAPE["right"]

HEX_NARRATION = "F3F9F5"
WORDS_PER_MINUTE = 140


# ---------------------------------------------------------------------------------------------
# Source
# ---------------------------------------------------------------------------------------------

def read_guide(source: Path) -> str:
    lines = source.read_text(encoding="utf-8").split("\n")
    try:
        start = next(i for i, line in enumerate(lines) if line.strip() == GUIDE_HEADING)
    except StopIteration:
        raise bm.BuildError(f"{source.name}: no '{GUIDE_HEADING}' heading")
    end = next((i for i in range(start + 1, len(lines)) if re.match(r"^##\s", lines[i])), len(lines))
    return "\n".join(lines[start + 1:end])


def shift_headings(text: str) -> str:
    """### becomes #, #### becomes ##, ##### becomes ###, outside code blocks."""
    out, in_code = [], False
    for line in text.split("\n"):
        if line.strip().startswith("```"):
            in_code = not in_code
        elif not in_code:
            m = re.match(r"^(#{3,5})(\s.*)$", line)
            if m:
                line = "#" * (len(m.group(1)) - 2) + m.group(2)
        out.append(line)
    return "\n".join(out)


def load_guide(source: Path) -> list[bm.Chapter]:
    text = shift_headings(read_guide(source))
    parts = re.split(r"(?m)^(?=# )", text)
    chapters = []
    for part in parts:
        if not part.startswith("# "):
            continue  # the lines before the first chapter are for readers of the plan
        label = str(len(chapters) + 1)
        chapter = bm.Chapter(Path(source.name), label, False)
        chapter.blocks = bm.parse_markdown(part, f"{source.name}, video guide chapter {label}")
        chapter.title = chapter.blocks[0].text
        chapter.anchor = chapter.blocks[0].anchor
        chapters.append(chapter)
    if not chapters:
        raise bm.BuildError(f"{source.name}: the video guide has no ### chapters")
    return chapters


def is_script(b: bm.Block) -> bool:
    return b.kind == "table" and bool(b.rows) and b.rows[0][0] == "Shot"


# ---------------------------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------------------------

def make_cover(version: str) -> Path:
    """A dark cover in the manual's style: the recipe editor picture shown as a paused video."""
    from PIL import Image, ImageDraw, ImageFilter
    import make_design_assets as art

    A4, mm, font = art.A4, art.mm, art.font
    img = art.gradient(A4, art.INK, art.INK_2)
    glow = Image.new("L", A4, 0)
    ImageDraw.Draw(glow).ellipse([mm(-20), mm(95), mm(230), mm(230)], fill=60)
    glow = glow.filter(ImageFilter.GaussianBlur(220))
    img.paste(Image.new("RGB", A4, (24, 70, 48)), (0, 0), glow)
    d = ImageDraw.Draw(img)
    left = mm(22)

    art.tracked(d, (left, mm(26)), "UNREAL ENGINE PLUGIN", font("seguisb.ttf", 10), art.GREEN, 9)
    d.text((left - 8, mm(33)), "FeelKit", font=font("seguisb.ttf", 76), fill=art.WHITE)
    d.text((left, mm(66)), "Video Guide", font=font("segoeuil.ttf", 30), fill=art.GREEN)
    d.text((left, mm(84)), "Recording, editing and publishing the tutorials", font=font("segoeui.ttf", 15), fill=art.WHITE)
    d.text((left, mm(92)), "and the listing videos, shot by shot.", font=font("segoeui.ttf", 15), fill=art.SOFT)

    if HERO.exists():
        shot = paused_player(Image.open(HERO).convert("RGB"))  # the whole editor window, as a screen recording shows it
        hero = art.framed(shot, A4[0] - left * 2 + 40)
        img.paste(hero, (left - 60 - 20, mm(108)), hero)
        d = ImageDraw.Draw(img)

    y = mm(232)
    col = (A4[0] - left * 2) // 3
    pillars = [
        ("SETUP", "OBS, microphone, Windows", "and Unreal, on this machine."),
        ("SCRIPTS", "Four tutorials and two listing", "videos, with every word."),
        ("PUBLISHING", "Editing, captions, loudness,", "YouTube and the Fab gallery."),
    ]
    for i, (title, line1, line2) in enumerate(pillars):
        x = left + i * col
        d.rectangle([x, y, x + mm(10), y + 6], fill=art.GREEN)
        art.tracked(d, (x, y + mm(5)), title, font("seguisb.ttf", 9.5), art.WHITE, 6)
        d.text((x, y + mm(12)), line1, font=font("segoeui.ttf", 10.5), fill=art.SOFT)
        d.text((x, y + mm(17.5)), line2, font=font("segoeui.ttf", 10.5), fill=art.SOFT)

    d.line([(left, mm(270)), (A4[0] - left, mm(270))], fill=(48, 54, 61), width=3)
    meta = f"FeelKit {version}        Video Guide        Billo"
    d.text((left, mm(275)), meta, font=font("segoeui.ttf", 9), fill=art.DIM)
    COVER.parent.mkdir(parents=True, exist_ok=True)
    img.save(COVER, dpi=(art.DPI, art.DPI), optimize=True)
    return COVER


def paused_player(shot):
    """The picture with a video player's controls: a round play button and a progress bar with a time."""
    from PIL import Image, ImageDraw, ImageFont
    import make_design_assets as art

    out = shot.convert("RGBA")
    layer = Image.new("RGBA", out.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    w, h = out.size
    s = w / 1400  # sizes below are for a picture 1400 px wide
    bar_h = int(64 * s)
    d.rectangle([0, h - bar_h, w, h], fill=(10, 12, 16, 225))
    track_y = h - bar_h + int(14 * s)
    x0, x1 = int(24 * s), w - int(24 * s)
    d.rounded_rectangle([x0, track_y, x1, track_y + int(6 * s)], int(3 * s), fill=(90, 98, 110, 255))
    played = x0 + int((x1 - x0) * 0.47)
    d.rounded_rectangle([x0, track_y, played, track_y + int(6 * s)], int(3 * s), fill=art.GREEN + (255,))
    knob = int(9 * s)
    d.ellipse([played - knob, track_y + int(3 * s) - knob, played + knob, track_y + int(3 * s) + knob], fill=(245, 247, 250, 255))
    small = ImageFont.truetype("C:/Windows/Fonts/segoeui.ttf", int(22 * s))
    d.polygon([(int(28 * s), h - int(36 * s)), (int(28 * s), h - int(14 * s)), (int(46 * s), h - int(25 * s))], fill=(245, 247, 250, 255))
    d.text((int(64 * s), h - int(40 * s)), "2:50 / 6:00", font=small, fill=(201, 209, 217, 255))
    cx, cy, r = int(w * 0.36), int(h * 0.38), int(62 * s)
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(10, 12, 16, 150), outline=(245, 247, 250, 230), width=max(2, int(4 * s)))
    d.polygon([(cx - int(18 * s), cy - int(30 * s)), (cx - int(18 * s), cy + int(30 * s)), (cx + int(32 * s), cy)], fill=(245, 247, 250, 240))
    return Image.alpha_composite(out, layer).convert("RGB")


# ---------------------------------------------------------------------------------------------
# Document
# ---------------------------------------------------------------------------------------------

class GuideBuilder(bm.Builder):
    """The manual's builder with its own title page, no part pages, and landscape pages for the scripts."""

    def __init__(self, version: str, release: bool):
        super().__init__(version, release, sample=False)
        self.landscape = False
        self.script_stats: list[tuple[str, int, float, list[str]]] = []

    # --- sections ---------------------------------------------------------------------------

    @staticmethod
    def geometry(section, landscape: bool) -> None:
        g = LANDSCAPE if landscape else PORTRAIT
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width = Cm(g["width"])
        section.page_height = Cm(g["height"])
        section.left_margin = Cm(g["left"])
        section.right_margin = Cm(g["right"])
        section.top_margin = Cm(g["top"])
        section.bottom_margin = Cm(g["bottom"])
        section.header_distance = Cm(g["header"])
        section.footer_distance = Cm(g["footer"])

    def running_header(self, section, width_cm: float) -> None:
        section.header.is_linked_to_previous = False
        header = section.header.paragraphs[0]
        header.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
        header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
        header.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)
        header.add_run(TITLE)
        header.add_run("\t")
        bm.add_field(header, 'STYLEREF "Heading 1"', "")
        for r in header.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = bm.COLOUR_MUTED
        hppr = header._p.get_or_add_pPr()
        hbdr = OxmlElement("w:pBdr")
        hbottom = OxmlElement("w:bottom")
        hbottom.set(qn("w:val"), "single")
        hbottom.set(qn("w:sz"), "4")
        hbottom.set(qn("w:space"), "6")
        hbottom.set(qn("w:color"), bm.HEX_RULE)
        hbdr.append(hbottom)
        hppr.insert(0, hbdr)
        section.footer.is_linked_to_previous = False
        self.page_numbers(section.footer.paragraphs[0])

    def new_section(self, landscape: bool = False, opener: bool = True) -> None:
        """Every section gets its own running header sized to its page, so portrait and landscape pages can alternate."""
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.geometry(section, landscape)
        section.different_first_page_header_footer = opener
        self.running_header(section, LANDSCAPE_TEXT if landscape else PORTRAIT_TEXT)
        self.header_ready = True
        bm.TEXT_WIDTH_CM = LANDSCAPE_TEXT if landscape else PORTRAIT_TEXT

    def start_landscape(self) -> None:
        self.new_section(landscape=True, opener=False)
        self.landscape = True

    def end_landscape(self) -> None:
        self.new_section(landscape=False, opener=False)
        self.landscape = False

    # --- title page -------------------------------------------------------------------------

    def title_page(self) -> None:
        doc = self.doc
        self.float_picture(make_cover(self.version), 21.0)
        contents = doc.add_paragraph("Contents", style="Contents Title")
        contents.paragraph_format.page_break_before = True
        toc = doc.add_paragraph()
        bm.add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Update fields (F9) to fill in the contents.")
        key = doc.add_paragraph()
        key.paragraph_format.space_before = Pt(20)
        self.add_pro_label(key, *bm.LABEL_SIZES["key"], gap=False)
        k = key.add_run("\u2003A video or shot that shows FeelKit Pro. Everything without this label applies to Lite and Pro.")
        k.font.size = Pt(8.5)
        k.font.color.rgb = bm.COLOUR_MUTED
        if not self.release:
            d = doc.add_paragraph().add_run(f"Draft, {dt.date.today().strftime('%d %B %Y').lstrip('0')}")
            d.font.size = Pt(8)
            d.font.color.rgb = bm.COLOUR_MUTED

    # --- script tables ----------------------------------------------------------------------

    def table(self, b: bm.Block) -> None:
        super().table(b)
        if not is_script(b):
            return
        table = self.doc.tables[-1]
        header = b.rows[0]
        narration = header.index("Narration") if "Narration" in header else -1
        for row in table.rows[1:]:
            cells = row.cells
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = bm.COLOUR_ACCENT
            for run in cells[1].paragraphs[0].runs:
                run.font.color.rgb = bm.COLOUR_MUTED
                run.text = run.text.replace("-", "‑")  # keep "0:00-0:15" on one line
            if narration >= 0:
                self.shade_cell(cells[narration], HEX_NARRATION)
                for run in cells[narration].paragraphs[0].runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = bm.COLOUR_HEADING
        self.record_stats(b)

    def record_stats(self, b: bm.Block) -> None:
        header = b.rows[0]
        if "Narration" not in header or "Time" not in header:
            return
        n_col, t_col = header.index("Narration"), header.index("Time")
        words, seconds, tight = 0, 0.0, []
        for row in b.rows[1:]:
            m = re.match(r"^(\d+):(\d\d)-(\d+):(\d\d)$", row[t_col].strip())
            if not m:
                raise bm.BuildError(f"{self.chapter.title}: shot {row[0]} time '{row[t_col]}' is not m:ss-m:ss")
            a = int(m.group(1)) * 60 + int(m.group(2))
            z = int(m.group(3)) * 60 + int(m.group(4))
            text = row[n_col].strip()
            count = 0 if text in ("", "None") else len(re.findall(r"[A-Za-z0-9][\w'.]*", text))
            words += count
            seconds = max(seconds, z)
            span = z - a
            if span <= 0:
                raise bm.BuildError(f"{self.chapter.title}: shot {row[0]} has no duration")
            if count and count / span * 60 > WORDS_PER_MINUTE * 1.15:
                tight.append(f"shot {row[0]}: {count} words in {span} s ({count / span * 60:.0f} per minute)")
        self.script_stats.append((self.chapter.title, words, seconds, tight))

    # --- whole document ---------------------------------------------------------------------

    def build_guide(self, chapters: list[bm.Chapter]) -> None:
        self.targets = bm.number_headings(chapters)
        self.planned = {}
        self.title_page()
        for ch in chapters:
            self.chapter = ch
            self.figure_in_chapter = 0
            self.last_heading = None
            blocks = ch.blocks
            for i, b in enumerate(blocks):
                goes_landscape = self.goes_landscape(blocks, i)
                if b.kind == "heading" and b.level == 1:
                    self.landscape = False  # the chapter opener starts its own portrait section
                elif goes_landscape and not self.landscape:
                    self.start_landscape()
                elif self.landscape and not goes_landscape:
                    self.end_landscape()
                self.dispatch(b)

    @staticmethod
    def goes_landscape(blocks: list[bm.Block], i: int) -> bool:
        """A script, the table right before it (its setup), and the heading above either go on the landscape page."""
        def at(k: int):
            return blocks[k] if k < len(blocks) else None
        b = blocks[i]
        if is_script(b):
            return True
        if b.kind == "table":
            return at(i + 1) is not None and is_script(at(i + 1))
        if b.kind == "heading" and b.level >= 2:
            nxt, after = at(i + 1), at(i + 2)
            return nxt is not None and (is_script(nxt) or (nxt.kind == "table" and after is not None and is_script(after)))
        return False

    def dispatch(self, b: bm.Block) -> None:
        handlers = {
            "heading": self.heading, "table": self.table, "list": self.list_block, "shot": self.figure,
            "video": self.video,
        }
        if b.kind in handlers:
            handlers[b.kind](b)
        elif b.kind == "para":
            self.paragraph(b.text)
        elif b.kind == "code":
            self.code(b.text)
        elif b.kind == "note":
            self.note(b.text)
        elif b.kind == "edition":
            self.edition(b.text)
        elif b.kind == "draft":
            self.draft(b.text)


# ---------------------------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------------------------

def lint(source: Path) -> list[str]:
    """The manual's writing rules on every line of the guide (code blocks and file names excepted)."""
    import lint_manual
    report = lint_manual.Report(release=False)
    lines = source.read_text(encoding="utf-8").split("\n")
    start = next(i for i, line in enumerate(lines) if line.strip() == GUIDE_HEADING)
    in_code = False
    for n in range(start, len(lines)):
        line = lines[n]
        if n > start and re.match(r"^##\s", line):
            break
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code or line.strip().startswith("{widths"):
            continue
        lint_manual.check_text_line(report, f"{source.name}:{n + 1}", line, prose=True)
    return report.errors + report.warnings


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Build the FeelKit Video Guide.")
    parser.add_argument("--pdf", action="store_true", help="update fields in Word and export a PDF next to the .docx")
    parser.add_argument("--release", action="store_true", help="no draft line on the title page")
    parser.add_argument("--source", type=Path, default=PLAN, help="file that holds the guide (default: the Manual Plan)")
    args = parser.parse_args()
    version = bm.plugin_version()
    bm.TITLE = TITLE  # the package properties written by build_manual.finalize_package
    try:
        chapters = load_guide(args.source)
        builder = GuideBuilder(version, args.release)
        builder.build_guide(chapters)
    except bm.BuildError as error:
        print(f"Build stopped: {error}", file=sys.stderr)
        return 1
    builder.doc.core_properties.author = bm.AUTHOR
    builder.doc.core_properties.title = TITLE
    DOCX.parent.mkdir(parents=True, exist_ok=True)
    builder.doc.save(DOCX)
    bm.finalize_package(DOCX, version, keep_app=False)
    print(f"wrote {DOCX}")
    if args.pdf:
        pdf = DOCX.with_suffix(".pdf")
        bm.export_with_word(DOCX.resolve(), pdf.resolve())
        bm.finalize_package(DOCX, version, keep_app=True)
        print(f"wrote {pdf}")
    print(f"  chapters: {len(chapters)}")
    for title, words, seconds, tight in builder.script_stats:
        need = words / WORDS_PER_MINUTE * 60
        print(f"  {title[:48]:48} {words:5} words, {need / 60:4.1f} min of speech in {seconds / 60:4.1f} min")
        for t in tight:
            print(f"      tight: {t}")
    problems = lint(args.source)
    for p in problems:
        print("lint:", p)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
